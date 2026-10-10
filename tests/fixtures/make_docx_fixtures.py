#!/usr/bin/env python3
"""Generate the committed .docx fixtures for extract_docx_style tests.

Run from the repo root:
    conda run -n py312 python tests/fixtures/make_docx_fixtures.py
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from pathlib import Path

OUT = Path(__file__).parent


def set_run(run, latin="Times New Roman", east="华文宋体", size=12,
            bold=False, italic=False, color=None, underline=False):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if underline:
        run.font.underline = True


def add_heading(doc, text, level, size, bold=True, color=None):
    h = doc.add_heading(text, level=level)
    set_run(h.runs[0], size=size, bold=bold, color=color)
    return h


def add_body(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=24,
             space_after=12, line_spacing=1.5, bold=False, italic=False,
             color=None, underline=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Pt(indent)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    r = p.add_run(text)
    set_run(r, size=12, bold=bold, italic=italic, color=color, underline=underline)
    return p


def mixed_numbering():
    doc = Document()
    doc.core_properties.title = "混合编号样例文档"
    add_heading(doc, "1. 总则", 1, size=22, color="1F4E79")
    add_heading(doc, "1.1 适用范围", 2, size=16)
    add_body(doc, "本文件规定了某事项的适用范围、基本原则与编号约定。")
    add_body(doc, "第二段正文，用于覆盖正文聚合统计。")
    add_heading(doc, "（一）分类", 2, size=16)
    add_body(doc, "本章节使用中文序号编号。", italic=True)
    add_heading(doc, "1.1.1 细则", 3, size=14)
    add_body(doc, "细则正文内容。", bold=True, color="C00000", underline=True)
    doc.save(OUT / "mixed_numbering.docx")


def dec_numbering():
    doc = Document()
    doc.core_properties.title = "十进制编号样例文档"
    add_heading(doc, "1. 总则", 1, size=22, color="1F4E79")
    add_heading(doc, "1.1 适用范围", 2, size=16)
    add_body(doc, "本文件规定十进制编号层级映射。")
    add_heading(doc, "1.1.1 细则", 3, size=14)
    add_body(doc, "细则正文。")
    add_heading(doc, "1.1.2 补充", 3, size=14)
    add_body(doc, "补充正文。")
    doc.save(OUT / "dec_numbering.docx")


if __name__ == "__main__":
    mixed_numbering()
    dec_numbering()
    print("fixtures generated:", OUT)
