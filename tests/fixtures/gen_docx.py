#!/usr/bin/env python3
"""Fixture: a safe, side-effect-free python-docx generator.

Used to exercise extract_docx_style.py's "input is a .py script" mode. Writes
only a single out.docx into the current working directory (a temp dir created
by the extractor). No network, no destructive file operations.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def set_run(run, latin="Times New Roman", east="华文宋体", size=12, bold=False):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.font.bold = bold


def add_heading(doc, text, level, size):
    h = doc.add_heading(text, level=level)
    set_run(h.runs[0], size=size, bold=True)


def main():
    doc = Document()
    doc.core_properties.title = "脚本生成文档"
    add_heading(doc, "1. 概述", 1, size=20)
    add_heading(doc, "1.1 步骤", 2, size=15)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Pt(24)
    pf.space_after = Pt(10)
    pf.line_spacing = 1.5
    r = p.add_run("由 py 生成脚本创建的正文段落。")
    set_run(r)
    doc.save("out.docx")


if __name__ == "__main__":
    main()
